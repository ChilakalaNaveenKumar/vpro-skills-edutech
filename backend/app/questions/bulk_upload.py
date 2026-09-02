"""Excel bulk-upload for questions (2026-08-31): parses a .xlsx workbook
into validated rows an admin router endpoint can turn into Question/
QuestionOption rows, so many questions can be added to a topic in one
upload instead of one form submission per question.

Kept as its own module (not schemas.py/router.py) since this is file
parsing and row-level validation for one specific input format, not a
Pydantic request/response shape or a route.
"""

from dataclasses import dataclass
from io import BytesIO

import openpyxl
from openpyxl.workbook import Workbook

from app.core.enums import EntityStatus

# Only the first 6 columns are actually required to be recognizable (by
# their leading word(s) - see the lenient header check below); the 7th
# (Status) is optional, so its exact wording isn't worth being strict
# about matching one of these header cells.
EXPECTED_HEADERS = [
    "Question Text",
    "Option A",
    "Option B",
    "Option C",
    "Option D",
    "Correct Option (A/B/C/D)",
    "Status (Active/Inactive) - optional, defaults to Active",
]

# Not a real product limit - just a sanity cap against an accidental
#10,000-row paste or a malicious upload tying up a request for too long.
MAX_ROWS = 500


@dataclass
class ParsedQuestionRow:
    question_text: str
    option_texts: dict[str, str]  # {"A": "...", "B": "...", "C": "...", "D": "..."}
    correct_label: str
    status: EntityStatus


@dataclass
class RowError:
    row: int  # 1-based, matching the row number the admin sees in Excel
    message: str


def build_template_workbook() -> Workbook:
    """The downloadable starting point for a bulk upload - exact column
    order/wording the parser below expects, plus two filled-in example
    rows so the format is obvious without reading a separate instructions
    doc. Regenerated fresh on every request (cheap, tiny) rather than
    served as a static file, so it can never drift from what the parser
    actually accepts.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Questions"
    ws.append(EXPECTED_HEADERS)
    ws.append(["What is 2 + 2?", "3", "4", "5", "22", "B", "Active"])
    ws.append(["Which keyword declares a constant in Java?", "var", "let", "final", "const", "C", ""])
    for column_cells in ws.columns:
        length = max(len(str(cell.value)) for cell in column_cells)
        ws.column_dimensions[column_cells[0].column_letter].width = min(max(length + 2, 12), 55)
    return wb


def _header_key(text: str) -> str:
    # "Correct Option (A/B/C/D)" -> "correct option"; "Status (Active/
    # Inactive) - optional, defaults to Active" -> "status" - strips
    # anything from the first "(" or " -" onward, so the parenthetical/
    # trailing hint text an admin might tweak or drop doesn't break the
    # match.
    return text.strip().lower().split(" (")[0].split(" -")[0]


def parse_questions_workbook(file_bytes: bytes) -> tuple[list[ParsedQuestionRow], list[RowError]]:
    """Returns (valid_rows, row_errors). Raises ValueError for a
    file-level problem (not a valid .xlsx, empty, wrong columns, too many
    rows) - the caller turns that into a 400 with the message as-is,
    since there's nothing row-specific to report. A wholly blank row
    (all six required cells empty) is silently skipped, not reported as
    an error - trailing blank rows are extremely common in hand-edited
    spreadsheets and aren't a mistake worth surfacing.
    """
    try:
        wb = openpyxl.load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(
            "Could not read this file as an Excel (.xlsx) workbook. Please use the template."
        ) from exc

    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("The file is empty.")

    header = [str(cell) if cell is not None else "" for cell in rows[0]]
    expected_keys = [_header_key(h) for h in EXPECTED_HEADERS[:6]]
    actual_keys = [_header_key(h) for h in header[:6]]
    if actual_keys != expected_keys:
        raise ValueError(
            "This file's columns don't match the expected template. Please download the "
            "template and add your questions to it, rather than starting from a blank sheet."
        )

    data_rows = rows[1:]
    if len(data_rows) > MAX_ROWS:
        raise ValueError(
            f"Too many rows ({len(data_rows)}). Please upload at most {MAX_ROWS} questions at a time."
        )

    parsed: list[ParsedQuestionRow] = []
    errors: list[RowError] = []

    for offset, row in enumerate(data_rows):
        row_number = offset + 2  # 1 for header + 1 to make this 1-based
        cells = list(row) + [None] * (7 - len(row))  # pad a short row (trailing columns omitted)
        question_text_raw, opt_a, opt_b, opt_c, opt_d, correct_raw, status_raw = cells[:7]

        required_cells = (question_text_raw, opt_a, opt_b, opt_c, opt_d, correct_raw)
        if all(c is None or str(c).strip() == "" for c in required_cells):
            continue

        row_errors: list[str] = []

        question_text = str(question_text_raw).strip() if question_text_raw is not None else ""
        if not question_text:
            row_errors.append("Question text is required.")

        option_texts: dict[str, str] = {}
        for label, value in (("A", opt_a), ("B", opt_b), ("C", opt_c), ("D", opt_d)):
            text = str(value).strip() if value is not None else ""
            if not text:
                row_errors.append(f"Option {label} is required.")
            option_texts[label] = text

        correct_label = str(correct_raw).strip().upper() if correct_raw is not None else ""
        if correct_label not in ("A", "B", "C", "D"):
            row_errors.append('"Correct Option" must be A, B, C, or D.')

        status_text = str(status_raw).strip().lower() if status_raw is not None else ""
        if status_text in ("", "active"):
            row_status = EntityStatus.ACTIVE
        elif status_text == "inactive":
            row_status = EntityStatus.INACTIVE
        else:
            row_errors.append('"Status" must be Active, Inactive, or left blank.')
            row_status = EntityStatus.ACTIVE  # unused - row is rejected below regardless

        if row_errors:
            errors.append(RowError(row=row_number, message=" ".join(row_errors)))
            continue

        parsed.append(
            ParsedQuestionRow(
                question_text=question_text,
                option_texts=option_texts,
                correct_label=correct_label,
                status=row_status,
            )
        )

    return parsed, errors
